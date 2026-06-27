from pathlib import Path

import docx
import fitz
import pptx
import pytesseract
from PIL import Image


class TextExtractionService:
    @staticmethod
    def extract_text(file_path: Path) -> str:
        extension = file_path.suffix.lower()

        if extension == ".txt":
            return file_path.read_text(encoding="utf-8", errors="ignore")
        if extension == ".pdf":
            return TextExtractionService._extract_pdf(file_path)
        if extension == ".docx":
            return TextExtractionService._extract_docx(file_path)
        if extension in {".ppt", ".pptx"}:
            return TextExtractionService._extract_pptx(file_path)
        if extension in {".png", ".jpg", ".jpeg"}:
            return TextExtractionService._extract_image(file_path)

        raise ValueError(f"Unsupported file type: {extension}")

    @staticmethod
    def _extract_pdf(file_path: Path) -> str:
        with fitz.open(file_path) as pdf:
            return "\n".join(page.get_text("text") for page in pdf)

    @staticmethod
    def _extract_docx(file_path: Path) -> str:
        document = docx.Document(file_path)
        return "\n".join(p.text for p in document.paragraphs)

    @staticmethod
    def _extract_pptx(file_path: Path) -> str:
        presentation = pptx.Presentation(file_path)
        chunks = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                text = getattr(shape, "text", "")
                if text:
                    chunks.append(text)
        return "\n".join(chunks)

    @staticmethod
    def _extract_image(file_path: Path) -> str:
        return pytesseract.image_to_string(Image.open(file_path))
